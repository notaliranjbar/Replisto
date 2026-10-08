"""Product catalogue: import from Google Sheets / Excel / CSV, column mapping, and search.

Every product lives in Replisto's own catalogue (`products`), whatever way the owner taught it:
typed in the panel, a live Google Sheet (re-synced every few minutes), an Excel/CSV upload,
or something the AI understood from text/a document (services/teach.py). The bot looks products up through a tool call, so prices and stock
always come from the owner's data instead of the model's imagination.

A sheet or file import only replaces the rows that came from that same source; products added
any other way are never touched by it.
"""

import csv
import functools
import io
import re
import uuid
from dataclasses import dataclass
from decimal import Decimal
from typing import Any
from urllib.parse import parse_qs, urlparse

import httpx
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.core.text import normalize, parse_decimal, parse_int
from app.core.timeutil import as_utc, utcnow
from app.models import Product, ProductOrigin, ProductSource, ProductSourceKind

# Our canonical fields, in the order the panel shows them. `name` is the only required one.
FIELDS = ("name", "sku", "category", "color", "size", "price", "stock", "description", "post_url")
FIELD_LABELS = {
    "name": "Product name",
    "sku": "SKU",
    "category": "Category",
    "color": "Colour",
    "size": "Size",
    "price": "Price",
    "stock": "Stock",
    "description": "Description",
    "post_url": "Instagram post link",
}
# Header names recognised for each field, in English and Arabic.
_SYNONYMS = {
    "name": [
        "product name",
        "اسم المنتج",
        "اسم",
        "الاسم",
        "المنتج",
        "منتج",
        "name",
        "product",
        "title",
        "item",
    ],
    "sku": ["sku", "رمز المنتج", "كود المنتج", "الرمز", "كود", "رمز", "code", "id"],
    "category": ["category", "الفئه", "فئه", "التصنيف", "تصنيف", "النوع", "type"],
    "color": ["colour", "color", "اللون", "لون"],
    "size": ["size", "المقاس", "مقاس", "الحجم"],
    "price": ["price", "السعر", "سعر", "الثمن"],
    "stock": ["stock", "المخزون", "مخزون", "الكميه", "كميه", "qty", "quantity", "inventory"],
    "description": [
        "description",
        "الوصف",
        "وصف",
        "التفاصيل",
        "الخامه",
        "desc",
        "details",
        "material",
    ],
    "post_url": ["post link", "رابط المنشور", "منشور", "انستغرام", "post", "instagram"],
}
TEMPLATE_HEADERS = [FIELD_LABELS[f] for f in FIELDS]
TEMPLATE_EXAMPLE = [
    "Linen dress",
    "M-101",
    "Dresses",
    "Black",
    "M",
    "49.90",
    "3",
    "100% linen",
    "https://www.instagram.com/p/XXXXXXXXXXX/",
]

_SHEETS_HOST = "docs.google.com"
_DRIVE_HOST = "drive.google.com"
_XLSX_MAGIC = b"PK"

_OUT_OF_STOCK_WORDS = ("غير متوفر", "غير متاح", "نفد", "نفذ", "out", "sold")
_IN_STOCK_WORDS = ("متوفر", "متاح", "موجود", "yes", "available", "in stock")


class ProductSourceError(Exception):
    """User-facing explanation of why the product table could not be used."""


@dataclass
class Table:
    headers: list[str]
    rows: list[dict[str, str]]


# ---------------------------------------------------------------------------- fetching


def download_url(url: str) -> str:
    """Turn a Google Sheets/Drive share link into a direct download URL.

    Only Google hosts are accepted, so users can't make our server fetch arbitrary URLs.
    """
    parsed = urlparse(url.strip())
    if parsed.scheme not in ("http", "https"):
        raise ProductSourceError("Invalid link.")
    if parsed.hostname == _SHEETS_HOST:
        match = re.search(r"/spreadsheets/d/([\w-]+)", parsed.path)
        if not match:
            raise ProductSourceError("Invalid Google Sheets link.")
        gid = parse_qs(parsed.query).get("gid", [None])[0]
        if gid is None:
            gid_match = re.search(r"gid=(\d+)", parsed.fragment)
            gid = gid_match.group(1) if gid_match else "0"
        return f"https://{_SHEETS_HOST}/spreadsheets/d/{match.group(1)}/export?format=csv&gid={gid}"
    if parsed.hostname == _DRIVE_HOST:
        match = re.search(r"/file/d/([\w-]+)", parsed.path)
        file_id = match.group(1) if match else parse_qs(parsed.query).get("id", [None])[0]
        if not file_id:
            raise ProductSourceError("Invalid Google Drive link.")
        return f"https://{_DRIVE_HOST}/uc?export=download&id={file_id}"
    raise ProductSourceError("Only Google Sheets links or Google Drive files are supported.")


async def fetch_table(http: httpx.AsyncClient, url: str) -> Table:
    try:
        resp = await http.get(download_url(url), follow_redirects=True, timeout=30)
    except httpx.HTTPError as exc:
        raise ProductSourceError("Couldn't reach Google. Please try again in a moment.") from exc
    content_type = resp.headers.get("content-type", "")
    if resp.status_code in (401, 403, 404) or "text/html" in content_type:
        raise ProductSourceError(
            "The file isn't accessible. In Google Sheets, open Share and set "
            "“Anyone with the link” as Viewer."
        )
    if resp.is_error:
        raise ProductSourceError("Google returned an error. Please try again in a moment.")
    if len(resp.content) > get_settings().upload_max_bytes:
        raise ProductSourceError("The product file is too large.")
    return parse_table(resp.content, "xlsx" if resp.content[:2] == _XLSX_MAGIC else "csv")


# ---------------------------------------------------------------------------- parsing


def parse_table(data: bytes, kind: str) -> Table:
    """Parse CSV or XLSX bytes. The first row with 2+ filled cells is the header row."""
    raw_rows = _xlsx_rows(data) if kind == "xlsx" else _csv_rows(data)
    header_index = next((i for i, r in enumerate(raw_rows) if sum(1 for c in r if c) >= 2), None)
    if header_index is None:
        raise ProductSourceError("The table is empty, or its header row couldn't be found.")

    headers: list[str] = []
    for i, cell in enumerate(raw_rows[header_index]):
        name = cell or f"Column {i + 1}"
        while name in headers:
            name += " (duplicate)"
        headers.append(name)

    rows = []
    for raw in raw_rows[header_index + 1 :]:
        if not any(raw):
            continue
        rows.append({h: (raw[i] if i < len(raw) else "") for i, h in enumerate(headers)})
    limit = get_settings().product_row_limit
    if len(rows) > limit:
        raise ProductSourceError(f"The table has {len(rows)} rows, over the limit of {limit}.")
    return Table(headers=headers, rows=rows)


def _csv_rows(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "utf-16", "cp1256"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ProductSourceError("Couldn't read the CSV file's encoding. Save it as UTF-8.")
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text))]


def _xlsx_rows(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt files
        raise ProductSourceError("Couldn't read the Excel file.") from exc
    sheet = workbook.worksheets[0]
    rows = [[_cell_text(c) for c in row] for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return rows


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# ---------------------------------------------------------------------------- mapping


def suggest_mapping(headers: list[str]) -> dict[str, str]:
    """Guess which sheet column holds each of our fields: exact synonym first, then substring."""
    mapping: dict[str, str] = {}
    used: set[str] = set()
    normalized = {h: normalize(h) for h in headers}
    for exact in (True, False):
        for field in FIELDS:
            if field in mapping:
                continue
            for synonym in _SYNONYMS[field]:
                target = normalize(synonym)
                if not exact and len(target) < 3:
                    continue  # "id": substring matches would hit unrelated headers
                hit = next(
                    (
                        h
                        for h in headers
                        if h not in used
                        and (normalized[h] == target if exact else target in normalized[h])
                    ),
                    None,
                )
                if hit:
                    mapping[field] = hit
                    used.add(hit)
                    break
    return mapping


def validate_mapping(mapping: dict[str, str], headers: list[str]) -> None:
    unknown_fields = set(mapping) - set(FIELDS)
    if unknown_fields:
        raise ProductSourceError(f"Unknown field: {', '.join(sorted(unknown_fields))}")
    missing = [h for h in mapping.values() if h not in headers]
    if missing:
        raise ProductSourceError(f"These columns aren't in the table: {', '.join(missing)}")
    if "name" not in mapping:
        raise ProductSourceError("Choose the “Product name” column.")


def parse_stock(value: str) -> int | None:
    number = parse_int(value)
    if number is not None:
        return max(number, 0)
    text = normalize(value)
    if not text:
        return None
    if any(w in text for w in _OUT_OF_STOCK_WORDS):  # before "متوفر", which "غير متوفر" contains
        return 0
    if any(w in text for w in _IN_STOCK_WORDS):
        return 1  # available, count unknown
    return None


def instagram_url(value: str) -> str | None:
    """Keep only real Instagram post/reel links; anything else in that column is ignored."""
    value = value.strip()
    return (
        value[:500] if re.match(r"^https?://(www\.)?instagram\.com/(p|reel|tv)/", value) else None
    )


def _mapped_value(row: dict[str, str], mapping: dict[str, str], field: str) -> str:
    return row.get(mapping[field], "").strip() if field in mapping else ""


def rows_to_products(table: Table, mapping: dict[str, str]) -> list[dict[str, Any]]:
    mapped_headers = set(mapping.values())
    products = []
    for row in table.rows:
        get = functools.partial(_mapped_value, row, mapping)
        name = get("name")
        if not name:
            continue
        price = parse_decimal(get("price"))
        products.append(
            {
                "name": name[:255],
                "sku": get("sku")[:64] or None,
                "category": get("category")[:120] or None,
                "color": get("color")[:64] or None,
                "size": get("size")[:64] or None,
                "price": price,
                "stock": parse_stock(get("stock")) if "stock" in mapping else None,
                "description": get("description") or None,
                "post_url": instagram_url(get("post_url")),
                "extra": {h: v for h, v in row.items() if h not in mapped_headers and v},
            }
        )
    return products


# ---------------------------------------------------------------------------- storage


def origin_for(source: ProductSource) -> ProductOrigin:
    return (
        ProductOrigin.SHEET if source.kind == ProductSourceKind.GOOGLE_SHEET else ProductOrigin.FILE
    )


async def replace_products(
    db: AsyncSession, source: ProductSource, products: list[dict[str, Any]]
) -> int:
    """Replace the rows imported from `source` (other products of the shop stay)."""
    await db.execute(delete(Product).where(Product.source_id == source.id))
    origin = origin_for(source)
    db.add_all(
        Product(shop_id=source.shop_id, source_id=source.id, origin=origin, **p) for p in products
    )
    await db.commit()
    return len(products)


async def sync_source(
    db: AsyncSession, source: ProductSource, http: httpx.AsyncClient, *, table: Table | None = None
) -> int:
    """Re-read a Google Sheet source (or use an already fetched `table`).

    Errors are recorded on the source and re-raised; the previous products stay in place.
    """
    assert source.url
    try:
        if table is None:
            table = await fetch_table(http, source.url)
        validate_mapping(source.column_mapping, table.headers)
        count = await replace_products(db, source, rows_to_products(table, source.column_mapping))
    except ProductSourceError as exc:
        await db.rollback()
        source.last_error = str(exc)
        await db.commit()
        raise
    source.last_synced_at = utcnow()
    source.last_error = None
    await db.commit()
    return count


# ---------------------------------------------------------------------------- search


def _haystack(p: Product) -> tuple[str, str]:
    primary = normalize(f"{p.name} {p.sku or ''}")
    secondary = normalize(
        " ".join(
            filter(None, [p.category, p.description, *(str(v) for v in (p.extra or {}).values())])
        )
    )
    return primary, secondary


def _variant_matches(value: str | None, wanted: str | None) -> bool:
    if not wanted:
        return True
    if value is None:
        return False
    a, b = normalize(value), normalize(wanted)
    return a == b or a in b or b in a


async def search_products(
    db: AsyncSession,
    shop_id: uuid.UUID,
    query: str,
    *,
    color: str | None = None,
    size: str | None = None,
    limit: int | None = 8,
) -> list[Product]:
    # A shop has at most product_row_limit rows, so scoring in Python is fine for now.
    products = (await db.scalars(select(Product).where(Product.shop_id == shop_id))).all()
    tokens = [t for t in normalize(query).split() if len(t) >= 2 or t.isdigit()]
    scored: list[tuple[int, Product]] = []
    for product in products:
        if not (_variant_matches(product.color, color) and _variant_matches(product.size, size)):
            continue
        primary, secondary = _haystack(product)
        score = sum(2 if t in primary else 1 if t in secondary else 0 for t in tokens)
        if tokens and score == 0:
            continue
        scored.append((score, product))
    scored.sort(key=lambda item: (-item[0], item[1].name))
    return [p for _, p in (scored[:limit] if limit else scored)]


def product_for_bot(p: Product) -> dict[str, Any]:
    info: dict[str, Any] = {"name": p.name}
    for key in ("sku", "category", "color", "size", "description"):
        if getattr(p, key):
            info[key] = getattr(p, key)
    if p.price is not None:
        info["price"] = price_number(p.price)
    if p.stock is not None:
        info["in_stock"] = p.stock > 0
        info["stock_count"] = p.stock
    if p.extra:
        info["other"] = p.extra
    return info


async def sync_due_sources(db: AsyncSession, http: httpx.AsyncClient) -> int:
    """Re-sync every Google Sheet source not synced within the configured interval."""
    interval = get_settings().product_sync_interval_minutes
    sources = (
        await db.scalars(
            select(ProductSource).where(ProductSource.kind == ProductSourceKind.GOOGLE_SHEET)
        )
    ).all()
    synced = 0
    for source in sources:
        last = source.last_synced_at
        if last is not None and (utcnow() - as_utc(last)).total_seconds() < (interval - 0.5) * 60:
            continue
        try:
            await sync_source(db, source, http)
            synced += 1
        except ProductSourceError:
            continue  # already recorded on the source; the panel shows it
    return synced


def _in_stock(p: Product) -> bool:
    return p.stock is None or p.stock > 0


def price_number(price: Decimal) -> int | float:
    """A price as a JSON number: 1250 or 49.9."""
    return int(price) if price == price.to_integral_value() else float(price)


def _money(price: Decimal, currency: str) -> str:
    """'$1,250', '$49.90' for USD; '1,250 AED' for other currencies."""
    number = f"{int(price):,}" if price == price.to_integral_value() else f"{price:,.2f}"
    return f"${number}" if currency == "USD" else f"{number} {currency}"


def catalogue_overview(
    items: list[Product],
    max_lines: int = 80,
    post_refs: dict[str, str] | None = None,
    currency: str = "USD",
) -> str:
    """One line per product name, cheapest first: price range, in-stock colours and sizes.

    Gives the bot a picture of the whole shop for browsing questions ("what do you sell?",
    "what's the cheapest?") that a name search can't answer. Exact stock stays with the
    search tool.
    """
    groups: dict[str, list[Product]] = {}
    for p in items:
        groups.setdefault(p.name.strip(), []).append(p)

    def unique(values: list[str | None]) -> list[str]:
        return list(dict.fromkeys(v.strip() for v in values if v and v.strip()))

    rows: list[tuple[int, str]] = []
    for name, variants in groups.items():
        prices = sorted({v.price for v in variants if v.price is not None})
        available = [v for v in variants if _in_stock(v)]
        parts = [f"- {name}"]
        category = next((v.category for v in variants if v.category), None)
        if category:
            parts[0] += f" ({category})"
        if prices:
            price = (
                _money(prices[0], currency)
                if len(prices) == 1
                else f"{_money(prices[0], currency)} to {_money(prices[-1], currency)}"
            )
            parts.append(price)
        if not available:
            parts.append("out of stock")
        else:
            colours = unique([v.color for v in available])
            sizes = unique([v.size for v in available])
            if colours:
                parts.append("colours in stock: " + ", ".join(colours))
            if sizes:
                parts.append("sizes in stock: " + ", ".join(sizes))
        if post_refs and name in post_refs:
            parts.append(f"post {post_refs[name]}")
        rows.append((prices[0] if prices else 10**15, "; ".join(parts)))

    rows.sort(key=lambda row: row[0])
    lines = [line for _, line in rows]
    if len(lines) > max_lines:
        hidden = len(lines) - max_lines
        lines = lines[:max_lines] + [f"- ...and {hidden} more products (use search_products)"]
    return "\n".join(lines)
