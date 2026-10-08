"""Shop knowledge: the guided sections, file-to-text extraction, and the prompt text."""

import io
from dataclasses import dataclass

from app.models import KnowledgeEntry, KnowledgeSection
from app.services.llm import LLMClient, LLMError


@dataclass(frozen=True)
class SectionInfo:
    key: KnowledgeSection
    label: str
    hint: str
    placeholder: str
    # Single sections are one text box in the panel; multiple ones are a list of entries.
    multiple: bool


# Order here is the order the panel shows them and the order the bot reads them.
SECTIONS: list[SectionInfo] = [
    SectionInfo(
        KnowledgeSection.ABOUT,
        "About the shop",
        "What does your shop sell, and what should we know about it?",
        "e.g. Layla Boutique makes women's dresses and blouses, all sewn in our own studio "
        "since 2019.",
        False,
    ),
    SectionInfo(
        KnowledgeSection.SHIPPING,
        "Shipping",
        "Shipping methods, costs, delivery times, other cities and international shipping",
        "e.g. Express delivery in 3-5 business days for $8. Free on orders over $150.",
        False,
    ),
    SectionInfo(
        KnowledgeSection.PAYMENT,
        "Payment",
        "Payment methods and how to order",
        "e.g. Card on our website, or cash on delivery within the city.",
        False,
    ),
    SectionInfo(
        KnowledgeSection.RETURNS,
        "Returns and exchanges",
        "Conditions and time limits for exchanges and returns",
        "e.g. Size exchanges within 7 days of delivery, unworn and with tags attached.",
        False,
    ),
    SectionInfo(
        KnowledgeSection.CONTACT,
        "Contact and opening hours",
        "Website, phone, store address, support hours",
        "e.g. layla.shop, support 10am-8pm, closed on Fridays.",
        False,
    ),
    SectionInfo(
        KnowledgeSection.FAQ,
        "Frequently asked questions",
        "Questions customers often ask, and their answers",
        "Q: What's the fabric? / A: Premium linen",
        True,
    ),
    SectionInfo(
        KnowledgeSection.FREE_TEXT,
        "Other details",
        "Anything else the bot should know",
        "e.g. 15% off your first order with the code FIRST15",
        True,
    ),
    SectionInfo(
        KnowledgeSection.FILE,
        "Uploaded files",
        "Text extracted from your files. You can edit it.",
        "",
        True,
    ),
]
SECTION_BY_KEY = {s.key: s for s in SECTIONS}

TEXT_TYPES = {".txt", ".md", ".csv"}
IMAGE_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}
SUPPORTED_EXTENSIONS = sorted(TEXT_TYPES | {".pdf", ".docx"} | IMAGE_TYPES.keys())


class ExtractionError(Exception):
    """User-facing explanation of why a file could not be read."""


async def extract_text(filename: str, data: bytes, llm: LLMClient) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in TEXT_TYPES:
        text = _decode(data)
    elif ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    elif ext in IMAGE_TYPES:
        if not llm.configured:
            raise ExtractionError(
                "Reading images isn't available yet. Please type the text instead."
            )
        try:
            text = await llm.extract_text_from_image(data, IMAGE_TYPES[ext])
        except LLMError as exc:
            raise ExtractionError("Reading the image failed. Please try again.") from exc
    else:
        allowed = ", ".join(SUPPORTED_EXTENSIONS)
        raise ExtractionError(f"This file type isn't supported. Allowed formats: {allowed}")

    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        raise ExtractionError(
            "No text was found in this file. If it's a scanned PDF, upload photos of the pages."
        )
    return text


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1256"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Couldn't read the text file's encoding. Save it as UTF-8.")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except (PdfReadError, ValueError) as exc:
        raise ExtractionError("The PDF is damaged or password-protected.") from exc


def _docx_text(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types for bad files
        raise ExtractionError("Couldn't read the Word file.") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def total_chars(entries: list[KnowledgeEntry]) -> int:
    return sum(len(e.title) + len(e.content) for e in entries)


def render_for_prompt(entries: list[KnowledgeEntry]) -> str:
    """Group entries under section headings, in the canonical section order."""
    blocks: list[str] = []
    for section in SECTIONS:
        items = [e for e in entries if e.section == section.key and e.content.strip()]
        if not items:
            continue
        lines = [f"## {section.label}"]
        for item in items:
            if item.title:
                lines.append(f"### {item.title}")
            lines.append(item.content.strip())
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)
