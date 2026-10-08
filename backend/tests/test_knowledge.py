import io

import docx
import pytest
from httpx import AsyncClient

from app.core.config import get_settings

API = "/api/v1/knowledge"


async def test_sections_describe_the_form(client: AsyncClient, auth_headers: dict) -> None:
    body = (await client.get(f"{API}/sections", headers=auth_headers)).json()
    keys = [s["key"] for s in body["sections"]]
    assert keys[:2] == ["about", "shipping"]
    faq = next(s for s in body["sections"] if s["key"] == "faq")
    assert faq["multiple"] is True
    assert ".pdf" in body["supported_file_types"]


async def test_single_section_upsert_and_clear(client: AsyncClient, auth_headers: dict) -> None:
    url = f"{API}/single/shipping"
    first = await client.put(url, headers=auth_headers, json={"content": "3-day express post"})
    assert first.status_code == 200
    second = await client.put(url, headers=auth_headers, json={"content": "2-day courier"})
    assert second.json()["id"] == first.json()["id"]  # updated, not duplicated

    listing = (await client.get(API, headers=auth_headers)).json()
    assert [e["content"] for e in listing["entries"]] == ["2-day courier"]
    assert listing["chars_used"] == len("2-day courier")

    cleared = await client.put(url, headers=auth_headers, json={"content": "  "})
    assert cleared.json() is None
    assert (await client.get(API, headers=auth_headers)).json()["entries"] == []


async def test_single_sections_reject_create(client: AsyncClient, auth_headers: dict) -> None:
    resp = await client.post(API, headers=auth_headers, json={"section": "about", "content": "x"})
    assert resp.status_code == 400
    resp = await client.put(f"{API}/single/faq", headers=auth_headers, json={"content": "x"})
    assert resp.status_code == 400


async def test_faq_crud(client: AsyncClient, auth_headers: dict) -> None:
    created = await client.post(
        API,
        headers=auth_headers,
        json={"section": "faq", "title": "What's the fabric?", "content": "Premium linen"},
    )
    assert created.status_code == 201
    entry_id = created.json()["id"]

    patched = await client.patch(
        f"{API}/{entry_id}", headers=auth_headers, json={"content": "Premium linen"}
    )
    assert patched.json()["content"] == "Premium linen"
    assert patched.json()["title"] == "What's the fabric?"

    assert (await client.delete(f"{API}/{entry_id}", headers=auth_headers)).status_code == 204
    assert (await client.delete(f"{API}/{entry_id}", headers=auth_headers)).status_code == 404


async def test_entries_are_private_to_their_shop(client: AsyncClient, auth_headers: dict) -> None:
    created = await client.post(
        API, headers=auth_headers, json={"section": "faq", "content": "Secret"}
    )
    other = await client.post(
        "/api/v1/auth/register",
        json={"email": "b@shop.com", "password": "strongpass1", "shop_name": "B"},
    )
    other_headers = {"Authorization": f"Bearer {other.json()['access_token']}"}
    entry_id = created.json()["id"]
    assert (await client.get(API, headers=other_headers)).json()["entries"] == []
    resp = await client.patch(f"{API}/{entry_id}", headers=other_headers, json={"content": "x"})
    assert resp.status_code == 404
    assert (await client.delete(f"{API}/{entry_id}", headers=other_headers)).status_code == 404


async def test_char_limit(
    client: AsyncClient, auth_headers: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(get_settings(), "knowledge_char_limit", 20)
    ok = await client.post(API, headers=auth_headers, json={"section": "faq", "content": "a" * 15})
    assert ok.status_code == 201
    over = await client.post(API, headers=auth_headers, json={"section": "faq", "content": "b" * 6})
    assert over.status_code == 413


async def test_upload_text_and_docx(client: AsyncClient, auth_headers: dict) -> None:
    txt = await client.post(
        f"{API}/upload",
        headers=auth_headers,
        files={"file": ("info.txt", b"Free shipping over $200", "text/plain")},
    )
    assert txt.status_code == 201
    assert txt.json()["section"] == "file"
    assert txt.json()["source_filename"] == "info.txt"
    assert txt.json()["content"] == "Free shipping over $200"

    document = docx.Document()
    document.add_paragraph("Returns: within 7 days")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Size M"
    table.rows[0].cells[1].text = "length 105"
    buffer = io.BytesIO()
    document.save(buffer)
    doc = await client.post(
        f"{API}/upload",
        headers=auth_headers,
        files={"file": ("rules.docx", buffer.getvalue(), "application/octet-stream")},
    )
    assert doc.status_code == 201
    assert "Returns: within 7 days" in doc.json()["content"]
    assert "Size M | length 105" in doc.json()["content"]


async def test_upload_rejections(client: AsyncClient, auth_headers: dict) -> None:
    async def upload(name: str, data: bytes) -> tuple[int, str]:
        resp = await client.post(
            f"{API}/upload", headers=auth_headers, files={"file": (name, data, "x/y")}
        )
        return resp.status_code, resp.json().get("detail", "")

    assert (await upload("a.exe", b"MZ"))[0] == 422
    assert (await upload("broken.pdf", b"not a pdf"))[0] == 422
    assert (await upload("empty.txt", b"   "))[0] == 422
    status, detail = await upload("menu.jpg", b"\xff\xd8fake")  # no LLM configured
    assert status == 422 and "image" in detail


async def test_upload_image_uses_llm(
    client: AsyncClient, auth_headers: dict, fake_web, llm_on
) -> None:
    fake_web.llm_replies = [{"content": "Linen dress: $125"}]
    resp = await client.post(
        f"{API}/upload",
        headers=auth_headers,
        files={"file": ("prices.png", b"\x89PNGfake", "image/png")},
    )
    assert resp.status_code == 201
    assert resp.json()["content"] == "Linen dress: $125"
    content = fake_web.llm_requests[0]["messages"][0]["content"]
    assert content[1]["image_url"]["url"].startswith("data:image/png;base64,")


async def test_upload_size_limit(
    client: AsyncClient, auth_headers: dict, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(get_settings(), "upload_max_bytes", 10)
    resp = await client.post(
        f"{API}/upload", headers=auth_headers, files={"file": ("big.txt", b"x" * 11, "text/plain")}
    )
    assert resp.status_code == 413
